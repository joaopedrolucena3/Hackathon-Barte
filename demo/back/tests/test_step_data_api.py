"""Data desk contract: whole step tables, control totals, cross-step invariants and exports of the visible rows."""
import csv
import io
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.app import create_app
from src.chat import semantic
from src.store import ROOT


class StepDataHTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        base = Path(cls.temp.name)
        cls.app = create_app(base/'case.sqlite3', ROOT/'data'/'case', base/'runs', mail_private_dir=base/'private')
        cls.client_context = TestClient(cls.app, base_url='http://127.0.0.1')
        cls.client = cls.client_context.__enter__()
        response = cls.client.post('/api/runs', headers={'X-User': 'Teste da mesa de dados'}, json={'pipeline_id': 'fechamento-tarvo', 'competencia': '2026-09'})
        assert response.status_code == 202, response.text
        cls.run_id = response.json()['run_id']
        until = time.monotonic() + 60
        while time.monotonic() < until:
            cls.state = cls.client.get('/api/runs/' + cls.run_id).json()
            if cls.state['status'] in ('awaiting_approval', 'failed'): break
            time.sleep(.05)
        assert cls.state['status'] == 'awaiting_approval', cls.state.get('error')

    @classmethod
    def tearDownClass(cls):
        cls.client_context.__exit__(None, None, None)
        cls.temp.cleanup()

    def data(self, step, dataset=None, **headers):
        path = f'/api/runs/{self.run_id}/steps/{step}/data' + (f'?dataset={dataset}' if dataset else '')
        response = self.client.get(path, headers=headers)
        self.assertEqual(response.status_code, 200, response.text)
        return response

    def export(self, **body):
        return self.client.post(f'/api/runs/{self.run_id}/steps/s5/export', json={'dataset': 'checkout_calculado', **body})

    def test_whole_table_keeps_file_order_types_and_exact_totals(self):
        table = self.data('s5').json()
        self.assertEqual((table['dataset'], table['row_count'], len(table['rows']), table['step_number']), ('checkout_calculado', 1943, 1943, 5))
        names = [column['name'] for column in table['columns']]
        gross = table['columns'][names.index('gross')]
        self.assertEqual((gross['label'], gross['kind']), ('Bruto', 'money'))
        self.assertTrue(gross['description'])
        self.assertTrue(next(c for c in table['columns'] if c['name'] == '__exception_ids')['hidden'])
        self.assertEqual(table['rows'][0][names.index('id')], 'TX88410932')
        self.assertIsInstance(table['rows'][0][names.index('gross')], str)
        totals = {item['column']: item['value'] for item in table['totals']}
        self.assertEqual((totals['gross'], totals['net'], totals['comissao_brl']), ('7496228.56', '7264348.20', '207039.73'))

    def test_invariants_compare_steps_and_metrics(self):
        checks = {item['id']: item for item in self.data('s5').json()['invariants']}
        commission = checks['comissao_checkout']
        self.assertTrue(commission['ok'])
        self.assertEqual({side['value'] for side in commission['sides']}, {'207039.73'})
        self.assertEqual(sorted(side['step_number'] for side in commission['sides']), [5, 5, 8])
        self.assertTrue(checks['bruto_preservado']['ok'])
        ledger = {item['id']: item for item in self.data('s8', 'lancamentos').json()['invariants']}
        self.assertTrue(ledger['lote_balanceado']['ok'])
        self.assertEqual(ledger['lote_balanceado']['sides'][0]['value'], '23474915.80')

    def test_large_tables_are_gzipped_and_identical(self):
        plain = self.data('s5').json()
        compressed = self.data('s5', **{'Accept-Encoding': 'gzip'})
        self.assertEqual(compressed.headers.get('content-encoding'), 'gzip')
        self.assertEqual(compressed.json(), plain)

    def test_dataset_selection_and_errors(self):
        notes = self.data('s8', 'notas').json()
        self.assertEqual(notes['dataset'], 'notas')
        self.assertIn('lancamentos', [item['name'] for item in notes['datasets']])
        self.assertEqual(self.client.get(f'/api/runs/{self.run_id}/steps/s8/data?dataset=raw_orbe').json()['error']['code'], 'DATASET_NOT_FOUND')
        self.assertEqual(self.client.get(f'/api/runs/{self.run_id}/steps/s99/data').json()['error']['code'], 'STEP_NOT_FOUND')

    def test_semantic_layer_documents_every_real_column(self):
        undocumented = []
        for step in self.state['steps']:
            for output in step['outputs']:
                for column in semantic.describe(output['name'], output.get('schema', [])):
                    if not column['description'] or column['label'] == column['name'].replace('_', ' ').capitalize() and column['name'] not in semantic.load()['rotulos']:
                        undocumented.append((output['name'], column['name']))
        self.assertEqual(undocumented, [])

    def test_csv_export_follows_screen_order_with_pt_br_values(self):
        response = self.export(rids=[2, 0], columns=['id', 'gross', 'date', 'calculated'], formato='csv')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn('attachment', response.headers['content-disposition'])
        text = response.content.decode('utf-8')
        self.assertTrue(text.startswith('﻿'))
        rows = list(csv.reader(io.StringIO(text.lstrip('﻿')), delimiter=';'))
        self.assertEqual(rows[0], ['Linha no dataset', 'ID', 'Bruto', 'Data', 'Calculada'])
        table = self.data('s5').json()
        names = [column['name'] for column in table['columns']]
        third = table['rows'][2]
        self.assertEqual(rows[1][0], '3')
        self.assertEqual(rows[1][1], third[names.index('id')])
        self.assertEqual(rows[1][2], third[names.index('gross')].replace('.', ','))
        self.assertRegex(rows[1][3], r'^\d{2}/\d{2}/\d{4}$')
        self.assertEqual(rows[2][1], 'TX88410932')

    def test_xlsx_export_writes_exact_decimals_and_provenance(self):
        response = self.export(rids=[0, 1], columns=['id', 'gross', 'comissao_brl'], formato='xlsx', descricao='Bandeira = ELO')
        self.assertEqual(response.status_code, 200, response.text)
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            sheet_xml = archive.read('xl/worksheets/sheet1.xml').decode()
        self.assertIn('<v>608.88</v>', sheet_xml)
        book = load_workbook(io.BytesIO(response.content))
        data = book['checkout_calculado']
        self.assertEqual(data['C2'].number_format, '#,##0.00')
        about = {row[0]: row[1] for row in book['Sobre'].iter_rows(values_only=True)}
        self.assertEqual(about['Filtros e ordenação'], 'Bandeira = ELO')
        self.assertEqual(about['SHA-256 da tabela'], self.data('s5').json()['sha256'])

    def test_export_rejects_unknown_columns_and_rows(self):
        self.assertEqual(self.export(rids=[0], columns=['inexistente'], formato='csv').json()['error']['code'], 'UNKNOWN_COLUMN')
        self.assertEqual(self.export(rids=[1943], columns=['id'], formato='csv').json()['error']['code'], 'INVALID_ROW')
        self.assertEqual(self.export(rids=[0], columns=['id'], formato='pdf').status_code, 422)


class InvariantUnitTests(unittest.TestCase):
    def test_mismatch_is_reported_with_both_values(self):
        steps = [{'id': 's1', 'number': 1, 'type': 'ler_arquivos', 'title': 'Ler arquivos', 'metrics': {}, 'datasets': ['raw_bank']},
                 {'id': 's6', 'number': 6, 'type': 'conciliar_aurora', 'title': 'Conferir bancos', 'metrics': {}, 'datasets': ['bancos_classificados']}]
        rows = {'raw_bank': [{'amount': '10.10'}, {'amount': '0.20'}], 'bancos_classificados': [{'valor': '10.30'}]}
        check = {item['id']: item for item in semantic.invariants(steps, rows.__getitem__)}['extrato_bancos']
        self.assertTrue(check['ok'])
        self.assertEqual([side['value'] for side in check['sides']], ['10.30', '10.30'])
        rows['bancos_classificados'] = [{'valor': '10.31'}]
        check = {item['id']: item for item in semantic.invariants(steps, rows.__getitem__)}['extrato_bancos']
        self.assertFalse(check['ok'])
        self.assertEqual([side['value'] for side in check['sides']], ['10.30', '10.31'])

    def test_invariant_needs_every_side(self):
        steps = [{'id': 's1', 'number': 1, 'type': 'ler_arquivos', 'title': 'Ler arquivos', 'metrics': {}, 'datasets': ['raw_bank']}]
        self.assertNotIn('extrato_bancos', {item['id'] for item in semantic.invariants(steps, lambda name: [])})


if __name__ == '__main__':
    unittest.main()
