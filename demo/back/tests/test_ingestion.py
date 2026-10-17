"""Input contracts: source receipts are not economic-event identities."""
import csv
import hashlib
import io
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

import openpyxl

from src.importer import load_data
from src.inputs import (ORBE_COLUMNS, ACCOUNT_COLUMNS, canonical_operation,
                        discover_files, inspect_bytes, inspect_file, operation_fingerprint)

PACK = Path(__file__).resolve().parents[1] / 'data' / 'case'


def orbe_bytes(rows=None, headers=ORBE_COLUMNS):
    if rows is None:
        rows = [['T-1', '30/09/2026', '1075624862', 'BRIDGE', 'VENDA',
                 'CREDITO A VISTA', '1', 'VISA', '100,00', '2,49', '2,49', '97,51', '', '']]
    output = io.StringIO(newline='')
    writer = csv.writer(output, delimiter=';')
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue().encode('utf-8')


class StructuralInspectionTests(unittest.TestCase):
    def test_six_sources_are_recognized_with_arbitrary_names(self):
        expected = {'orbe': 1943, 'taxas': 25, 'ponte': 233, 'contas': 52, 'aurora': 84, 'meridiano': 100}
        originals = discover_files(PACK)
        self.assertEqual(len(originals), 6)
        for entry in originals:
            with self.subTest(role=entry['role']):
                result = inspect_bytes('fonte-sem-prefixo' + Path(entry['name']).suffix.upper(), Path(entry['path']).read_bytes())
                self.assertEqual(result['status'], 'valid')
                self.assertEqual(result['role'], entry['role'])
                self.assertEqual(result['row_count'], expected[entry['role']])

    def test_filename_is_not_evidence_of_layout(self):
        result = inspect_bytes('01_orbe_relatorio_comissoes.csv', b'hello;world\na;b\n')
        self.assertEqual(result['status'], 'invalid')
        self.assertIsNone(result['role'])
        self.assertTrue(result['reason'])
        self.assertEqual(result['sha256'], hashlib.sha256(b'hello;world\na;b\n').hexdigest())

    def test_empty_header_only_and_bad_layout_are_distinct(self):
        cases = [(b'', 'empty'), (b'\xef\xbb\xbf\r\n  ', 'empty'),
                 (orbe_bytes([]), 'header_only'), (b'id_transacao;valor_bruto\nT;10\n', 'invalid')]
        for raw, status in cases:
            with self.subTest(status=status):
                result = inspect_bytes('recebido.csv', raw)
                self.assertEqual(result['status'], status)
                self.assertEqual(result['validation']['status'], status)
                self.assertEqual(result['size'], len(raw))

    def test_selected_role_must_match_structure(self):
        self.assertEqual(inspect_bytes('qualquer.csv', orbe_bytes(), role='orbe')['status'], 'valid')
        result = inspect_bytes('qualquer.csv', orbe_bytes(), role='contas')
        self.assertEqual(result['status'], 'invalid')
        self.assertEqual(result['candidate_roles'], ['orbe'])
        self.assertEqual(inspect_bytes('qualquer.csv', orbe_bytes(), role='inventado')['status'], 'invalid')

    def test_ambiguous_layout_requires_explicit_validated_choice(self):
        headers = list(ORBE_COLUMNS) + list(ACCOUNT_COLUMNS)
        raw = orbe_bytes([[''] * len(headers)], headers)
        # A non-empty row lets this verify layout ambiguity separately from emptiness.
        raw = raw.replace(b';;;;;;;;;;;;;;;;', b'T;100;;;;;;;;;;;;;;', 1)
        result = inspect_bytes('duas-estruturas.csv', raw)
        self.assertEqual(result['status'], 'ambiguous')
        self.assertEqual(result['candidate_roles'], ['contas', 'orbe'])
        selected = inspect_bytes('duas-estruturas.csv', raw, role='orbe')
        self.assertEqual(selected['role'], 'orbe')
        self.assertEqual(selected['status'], 'valid')

    def test_duplicate_required_header_is_invalid(self):
        raw = orbe_bytes([], list(ORBE_COLUMNS) + ['id_transacao'])
        self.assertEqual(inspect_bytes('duplicado.csv', raw)['status'], 'invalid')

    def test_workbook_with_only_header_is_explicit(self):
        original = next(PACK.glob('02*.xlsx'))
        book = openpyxl.load_workbook(io.BytesIO(original.read_bytes()))
        sheet = book['Taxas vigentes']
        sheet.delete_rows(6, sheet.max_row)
        data = io.BytesIO()
        book.save(data)
        book.close()
        result = inspect_bytes('contratos.xlsx', data.getvalue())
        self.assertEqual((result['role'], result['status'], result['row_count']), ('taxas', 'header_only', 0))

    def test_corrupt_workbook_and_unreadable_file_are_manifested(self):
        result = inspect_bytes('contratos.xlsx', b'PKbroken')
        self.assertEqual(result['status'], 'invalid')
        with tempfile.TemporaryDirectory() as folder:
            missing = inspect_file(Path(folder) / 'missing.csv', role='orbe', file_id='receipt-9')
        self.assertEqual(missing['id'], 'receipt-9')
        self.assertEqual(missing['status'], 'invalid')

    def test_explicit_bank_metadata_cannot_silently_change_account_identity(self):
        source = next(PACK.glob('05b*.ofx')).read_bytes()
        for changed in (source.replace(b'<BANKID>0477', b'<BANKID>9999'),
                        source.replace(b'<ACCTID>90233-5', b'<ACCTID>00000-0')):
            with self.subTest(format='OFX'):
                self.assertEqual(inspect_bytes('qualquer.ofx', changed, role='meridiano')['status'], 'invalid')
        source = next(PACK.glob('05a*.csv')).read_bytes()
        for changed in (source.replace(b'Banco Aurora', b'Banco Desconhecido'),
                        source.replace(b'Conta: 4471-0', b'Conta: 9999-9')):
            with self.subTest(format='CSV'):
                self.assertEqual(inspect_bytes('qualquer.csv', changed, role='aurora')['status'], 'invalid')


class ReceiptImportTests(unittest.TestCase):
    def test_discovery_retains_every_data_file_and_ignores_documentation(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'arbitrario.csv').write_bytes(orbe_bytes())
            (root / 'vazio.csv').write_bytes(b'')
            (root / 'ruim.xlsx').write_bytes(b'broken workbook')
            (root / 'README.md').write_text('Documentação')
            (root / 'manual.docx').write_bytes(b'not an input')
            data = load_data(root)
        self.assertEqual({r['name'] for r in data['input_manifest']}, {'arbitrario.csv', 'vazio.csv', 'ruim.xlsx'})
        self.assertEqual(len(data['sources']), 3)
        self.assertEqual(len(data['checkout']), 1)
        types = {w['type'] for w in data['import_warnings']}
        self.assertTrue({'ARQUIVO_VAZIO', 'ARQUIVO_INVALIDO'}.issubset(types))

    def test_unknown_extensions_are_inspected_like_upload_and_never_silently_ignored(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'foo.bin').write_bytes(orbe_bytes())
            (root / 'incompativel.bin').write_bytes(b'not a financial source')
            (root / 'README.txt').write_text('Documentação')
            (root / '.DS_Store').write_bytes(b'OS metadata')
            (root / '~$editing.xlsx').write_bytes(b'temporary editor lock')
            data = load_data(root)
        manifest = {entry['name']: entry for entry in data['input_manifest']}
        self.assertEqual(set(manifest), {'foo.bin', 'incompativel.bin'})
        self.assertEqual((manifest['foo.bin']['role'], manifest['foo.bin']['status']), ('orbe', 'valid'))
        self.assertEqual(manifest['incompativel.bin']['status'], 'invalid')
        self.assertEqual(len(data['checkout']), 1)
        self.assertEqual(inspect_bytes('foo.bin', orbe_bytes())['status'], manifest['foo.bin']['status'])

    def test_complete_folder_uses_structure_even_when_all_extensions_are_arbitrary(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for index, entry in enumerate(discover_files(PACK)):
                (root / f'fonte-{index}.bin').write_bytes(Path(entry['path']).read_bytes())
            data = load_data(root)
        self.assertEqual(len(data['input_manifest']), 6)
        self.assertTrue(all(entry['status'] == 'valid' for entry in data['input_manifest']))
        self.assertFalse(data['import_warnings'])
        self.assertEqual((len(data['checkout']), len(data['bank']), len(data['remittances'])), (1943, 184, 233))

    def test_multiple_orbe_original_name_hash_and_physical_row(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            first = root / 'a.csv'
            second = root / 'b.csv'
            first.write_bytes(orbe_bytes())
            second.write_bytes(orbe_bytes().replace(b'T-1;', b'T-2;'))
            data = load_data(root)
        self.assertEqual(len(data['checkout']), 2)
        for row in data['checkout']:
            self.assertEqual(row['physical_row'], 2)
            entry = next(e for e in data['input_manifest'] if e['id'] == row['source_id'])
            self.assertEqual(row['source_file'], entry['name'])
            self.assertEqual(row['source_sha256'], entry['sha256'])
        self.assertEqual(data['ingestion_counts']['by_role']['orbe']['rows_received'], 2)
        self.assertEqual(data['ingestion_counts']['by_role']['orbe']['rows_unique'], 2)

    def test_same_name_different_receipts_never_overwrites_or_changes_identity(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            receipts = []
            for identity, raw in [('first', orbe_bytes()), ('second', orbe_bytes().replace(b'T-1;', b'T-2;'))]:
                directory = root / identity
                directory.mkdir()
                path = directory / 'same.csv'
                path.write_bytes(raw)
                receipts.append({'id': identity, 'name': 'same.csv', 'path': str(path), 'role': 'orbe',
                                 'sha256': hashlib.sha256(raw).hexdigest()})
            before = [Path(r['path']).read_bytes() for r in receipts]
            data = load_data(root, files=receipts)
            self.assertEqual([Path(r['path']).read_bytes() for r in receipts], before)
        self.assertEqual({r['source_id'] for r in data['checkout']}, {'first', 'second'})
        self.assertEqual({r['source_file'] for r in data['checkout']}, {'same.csv'})
        self.assertEqual(len({r['source_sha256'] for r in data['checkout']}), 2)
        self.assertEqual(len(data['input_manifest']), 2)

    def test_expected_hash_is_checked_again(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'a.csv'
            path.write_bytes(orbe_bytes())
            data = load_data(folder, [{'id': 'id', 'name': 'original.csv', 'path': str(path), 'role': 'orbe', 'sha256': '0' * 64}])
        self.assertFalse(data['checkout'])
        self.assertEqual(data['input_manifest'][0]['status'], 'invalid')
        self.assertIn('hash', data['input_manifest'][0]['reason'])

    def test_identical_single_role_receipts_keep_all_origins_and_parse_once(self):
        original = next(PACK.glob('05b*.ofx')).read_bytes()
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'first.ofx').write_bytes(original)
            (root / 'renamed.ofx').write_bytes(original)
            manifest = discover_files(root)
            for index, entry in enumerate(manifest):
                entry['receipt_id'] = 'receipt-' + str(index)
            data = load_data(root, manifest)
        self.assertEqual(len(data['bank']), 100)
        self.assertEqual(len(data['input_manifest']), 2)
        self.assertTrue(all(row['status'] == 'valid' for row in data['input_manifest']))
        for row in data['bank']:
            self.assertEqual({p['source_file'] for p in row['provenances']}, {'first.ofx', 'renamed.ofx'})
            self.assertEqual(len({p['source_id'] for p in row['provenances']}), 2)
            self.assertEqual({p['source_receipt_id'] for p in row['provenances']}, {'receipt-0', 'receipt-1'})
        counts = data['ingestion_counts']['by_role']['meridiano']
        self.assertEqual((counts['rows_received'], counts['rows_valid']), (200, 200))

    def test_distinct_single_role_sources_are_ambiguous_and_none_selected(self):
        original = next(PACK.glob('05b*.ofx')).read_bytes()
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'first.ofx').write_bytes(original)
            (root / 'second.ofx').write_bytes(original + b'\n')
            data = load_data(root)
        self.assertFalse(data['bank'])
        self.assertEqual([row['status'] for row in data['input_manifest']], ['ambiguous', 'ambiguous'])
        self.assertEqual(sum(w['type'] == 'ARQUIVO_AMBIGUO' for w in data['import_warnings']), 2)

    def test_invalid_values_have_received_valid_discarded_and_physical_evidence(self):
        original = orbe_bytes()
        bad = original.decode().splitlines()[1].replace('T-1;', 'T-BAD;').replace('100,00', 'NaN')
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'mes.csv'
            path.write_bytes(original + (bad + '\r\n').encode())
            data = load_data(folder)
        entry = data['input_manifest'][0]
        self.assertEqual((entry['rows_received'], entry['rows_valid'], entry['rows_discarded']), (2, 1, 1))
        warning = next(w for w in data['import_warnings'] if w['type'] == 'LINHA_INVALIDA')
        self.assertEqual((warning['physical_row'], warning['source_ref'], warning['source_id']), (3, 'L3', entry['id']))

    def test_short_row_is_pending_not_silently_filled(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'mes.csv'
            path.write_bytes(orbe_bytes([]) + b'T-1;30/09/2026\n')
            data = load_data(folder)
        self.assertFalse(data['checkout'])
        self.assertEqual(data['input_manifest'][0]['rows_discarded'], 1)
        self.assertTrue(any(w['type'] == 'LINHA_INVALIDA' for w in data['import_warnings']))

    def test_multiline_csv_records_retain_starting_physical_lines(self):
        row = ['T-1', '30/09/2026', '1075624862', 'Nome\nComposto', 'VENDA', 'CREDITO A VISTA', '1',
               'VISA', '100,00', '2,49', '2,49', '97,51', '', '']
        second = list(row)
        second[0] = 'T-2'
        second[3] = 'Simples'
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'mes.csv'
            path.write_bytes(orbe_bytes([row, second]))
            data = load_data(folder)
        self.assertEqual([r['physical_row'] for r in data['checkout']], [2, 4])

    def test_receipt_counts_distinguish_duplicates_and_conflicting_ids(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for name, raw in [('one.csv', orbe_bytes()), ('same.csv', orbe_bytes()),
                              ('conflict.csv', orbe_bytes().replace(b'100,00', b'200,00'))]:
                (root / name).write_bytes(raw)
            data = load_data(root)
            (root / 'conflict.csv').unlink()
            duplicate = load_data(root)
        counts = data['ingestion_counts']['by_role']['orbe']
        self.assertEqual((counts['rows_received'], counts['rows_valid'], counts['rows_unique'], counts['rows_conflicting']), (3, 3, 0, 3))
        counts = duplicate['ingestion_counts']['by_role']['orbe']
        self.assertEqual((counts['rows_unique'], counts['rows_duplicate'], counts['rows_conflicting']), (1, 1, 0))
        self.assertEqual(len(duplicate['checkout']), 2, 'Dedup is owned by engine, while import preserves receipts')

    def test_contract_identity_does_not_depend_on_physical_order(self):
        original = next(PACK.glob('02*.xlsx'))
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / 'taxas.xlsx'
            source.write_bytes(original.read_bytes())
            before = load_data(root)
            book = openpyxl.load_workbook(io.BytesIO(original.read_bytes()))
            sheet = book['Taxas vigentes']
            rows = list(sheet.values)[5:30]
            for line, values in enumerate(reversed(rows), 6):
                for column, value in enumerate(values, 1):
                    sheet.cell(line, column).value = value
            book.save(source)
            book.close()
            after = load_data(root)
        keys = lambda data: sorted((r['logical_id'], r['revision_id'], r['rate_pct']) for r in data['rates'])
        self.assertEqual(keys(before), keys(after))
        self.assertNotEqual([r['source_ref'] for r in before['rates']], [r['source_ref'] for r in after['rates']])


class EconomicIdentityTests(unittest.TestCase):
    def test_one_installment_credit_matches_raw_history_and_normalized_current_row(self):
        raw = {'id': 'ORIGINAL-1', 'modality': 'CREDITO PARCELADO', 'installments': 1, 'gross': '100.00'}
        normalized = {**raw, 'modality': 'CREDITO A VISTA', 'installments': '1.0'}
        self.assertEqual(operation_fingerprint(raw), operation_fingerprint(normalized))
        self.assertNotEqual(operation_fingerprint(raw), operation_fingerprint({**raw, 'installments': 2}))
        self.assertNotEqual(operation_fingerprint(raw), operation_fingerprint({**raw, 'installments': None}))
        # Anticipation uses the first branch of the normaliser and keeps modality.
        self.assertNotEqual(operation_fingerprint({**raw, 'type': 'ANTECIPACAO'}),
                            operation_fingerprint({**normalized, 'type': 'ANTECIPACAO'}))

    def test_canonical_numeric_identity_never_rounds_to_decimal_context_precision(self):
        first = {'id': 'T', 'gross': '1.00000000000000000000000000001'}
        different = {'id': 'T', 'gross': '1.00000000000000000000000000002'}
        self.assertNotEqual(operation_fingerprint(first), operation_fingerprint(different))
        self.assertEqual(operation_fingerprint(first), operation_fingerprint({**first, 'gross': first['gross'] + '0'}))

    def test_numeric_equivalence_receipts_and_enrichment_do_not_change_identity(self):
        source = {'id': 'T', 'gross': '100.00', 'installments': 1, 'applied_rate_pct': '2.4900', 'net': '97.510',
                  'ec': '00123', 'brand': 'VISA', 'date': '2026-09-01', 'source_file': 'a.csv', 'physical_row': 2}
        other = {**source, 'gross': '1E+2', 'installments': '1.0', 'applied_rate_pct': '2.49', 'net': '97.51',
                 'source_file': 'b.csv', 'source_id': 'receipt-2', 'source_sha256': 'new', 'physical_row': 9,
                 'provenances': [source], 'original_evidence': source, 'original_resolution': 'resolved',
                 'accumulated_reversal_brl': '25', 'reversal_evidence': [source], '__exception_ids': ['e1']}
        self.assertEqual(operation_fingerprint(source), operation_fingerprint(other))
        self.assertEqual(canonical_operation(source)['ec'], '00123')

    def test_real_business_change_and_missing_invalid_amounts_remain_distinct(self):
        original = {'id': 'T', 'gross': '100.00', 'ec': '123'}
        fingerprints = {operation_fingerprint({**original, 'gross': value}) for value in ['100', '101', '', None, 'NaN', 'invalid']}
        self.assertEqual(len(fingerprints), 6)
        self.assertNotEqual(operation_fingerprint(original), operation_fingerprint({**original, 'ec': '124'}))


if __name__ == '__main__':
    unittest.main()
