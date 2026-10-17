"""Shared, content-based input inspection and operation identity.

Inspection never changes input bytes. It describes every receipt, including rejected
files; a recognised extension/name is not sufficient to accept a source.
"""
from __future__ import annotations

import csv
import email
import hashlib
import io
import json
import re
from decimal import Decimal
from email import policy
from pathlib import Path

import openpyxl

ROLES = ("orbe", "taxas", "ponte", "contas", "aurora", "meridiano")
DATA_EXTENSIONS = frozenset({".csv", ".xlsx", ".eml", ".ofx"})
DOCUMENT_EXTENSIONS = frozenset({".md", ".rst", ".pdf", ".doc", ".docx", ".rtf"})
DOCUMENT_NAMES = frozenset({"readme", "hipoteses", "leia-me", "leiame", "license", "licence", "agents", "changelog", "contributing"})
ORBE_COLUMNS = (
    "id_transacao", "data_movimento", "cod_estabelecimento", "nome_estabelecimento",
    "tipo_lancamento", "modalidade", "qtd_parcelas", "bandeira", "valor_bruto",
    "taxa_aplicada_pct", "valor_taxa", "valor_liquido", "id_referencia", "data_venda_original",
)
ACCOUNT_COLUMNS = ("COD_CONTA", "DESCRICAO", "NATUREZA", "ACEITA_LANCTO")
AURORA_COLUMNS = ("Data", "Lançamento", "Documento", "Valor (R$)", "Saldo (R$)")
RATE_COLUMNS = ("Lojista", "CNPJ", "EC Orbe", "Débito", "Crédito à vista", "Créd. 2-6x",
                "Créd. 7-12x", "Pix", "Antecipação", "Spread remessa", "Vigência", "Obs")
PONTE_COLUMNS = ("Nº Operação", "Data Contratação", "Data Liquidação", "CNPJ Cliente",
                 "Cliente", "Moeda", "Valor ME", "Taxa Cliente", "PTAX Ref.", "Valor R$",
                 "Spread R$", "Spread %", "Spread Estimado R$ (D-1)", "IOF R$",
                 "Status Compliance", "Observação")

# These columns describe a receipt/evidence, not an economic operation. The engine
# can retain them in trace details while comparing canonical business content.
PROVENANCE_FIELDS = frozenset({
    "source_file", "source_ref", "physical_row", "source_id", "source_sha256",
    "source_hash", "file_id", "file_sha256", "source_receipt_id", "receipt_id", "path", "sources", "origins",
    "provenances", "provenance", "observation", "__exception_ids", "history_run_id",
    "history_run_ids", "historical", "history_only", "source_run_id", "source_run_ids",
    "attachment_sha256", "original_evidence", "original_resolution", "accumulated_reversal_brl",
    "reversal_evidence",
})
NUMERIC_FIELDS = frozenset({
    "gross", "applied_rate_pct", "reported_fee", "net", "foreign_amount", "client_rate",
    "ptax_rate", "gross_brl", "spread_brl", "estimated_spread_brl", "iof_brl", "amount",
    "balance", "installments", "rate_pct",
})

# Identity evidence in the supplied case: Aurora CSV lines 1/4, and the
# Meridiano OFX BANKACCTFROM/BANKID and ACCTID fields. The accounting map routes
# these two accounts separately. A different explicit account must not inherit
# their identity merely because its statement uses the same layout.
BANK_IDENTITIES = {
    "aurora": {"institution": "Banco Aurora", "account": "44710"},
    "meridiano": {"bank_id": "477", "account": "902335"},
}


def canonical_operation(row):
    """Business comparison independent of receipt and decimal representation.

    Invalid and missing numbers remain distinct; validation belongs to the engine,
    so canonicalisation must never turn bad input into an amount of zero.
    """
    result = {}
    for key, value in row.items():
        if key in PROVENANCE_FIELDS or key.startswith("__"):
            continue
        if key in NUMERIC_FIELDS and value is not None and value != "":
            try:
                number = Decimal(str(value))
                if number.is_finite():
                    # normalize() rounds through the active Decimal context;
                    # formatting and trimming preserve every input digit.
                    value = "0" if not number else format(number, "f")
                    if "." in value:
                        value = value.rstrip("0").rstrip(".")
            except (ArithmeticError, ValueError):
                pass
        result[key] = value
    # Same existing normalisation as the Orbe pipeline: a one-installment
    # credit sale is cash credit. Historical raw evidence must compare equal
    # to that operation after the normalisation step.
    if (result.get("modality") == "CREDITO PARCELADO" and result.get("installments") == "1"
            and result.get("type") not in {"ANTECIPACAO", "REVERSAO ANTECIPACAO"}):
        result["modality"] = "CREDITO A VISTA"
    return result


def operation_fingerprint(row):
    return hashlib.sha256(json.dumps(canonical_operation(row), sort_keys=True,
                                   ensure_ascii=False, separators=(",", ":"),
                                   default=str).encode("utf-8")).hexdigest()


def decode_text(raw):
    try:
        return raw.decode("utf-8-sig"), "utf-8-sig"
    except UnicodeDecodeError:
        return raw.decode("cp1252"), "cp1252"


def csv_rows(text, delimiter=";"):
    """Yield physical *starting* line and parsed columns (including multiline CSV)."""
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    while True:
        line = reader.line_num + 1
        try:
            cells = next(reader)
        except StopIteration:
            return
        yield line, cells


def matches_header(cells, required):
    values = [str(v).strip() if v is not None else "" for v in cells]
    return all(values.count(column) == 1 for column in required)


def workbook_rows(book, sheet_name, required):
    """Return header and business rows, omitting only documented report footers."""
    if sheet_name not in book.sheetnames:
        return None, [], []
    headers, header_line, rows = None, None, []
    for line, cells in enumerate(book[sheet_name].values, 1):
        if headers is None and matches_header(cells, required):
            headers = [str(c).strip() if c is not None else "" for c in cells]
            header_line = line
            continue
        if headers is None or not any(c is not None and str(c).strip() for c in cells):
            continue
        row = dict(zip(headers, cells))
        if sheet_name == "Taxas vigentes":
            # One explanatory line below the report has no CNPJ and no rates.
            if str(row.get("Lojista") or "").startswith("*") and not row.get("CNPJ"):
                continue
        elif str(row.get("Nº Operação") or "").strip() in {"Total de operações:", "Spread total R$:"}:
            continue
        rows.append((line, row))
    return header_line, headers or [], rows


def ponte_attachments(raw):
    """Find structurally compatible workbooks, regardless of attachment name."""
    message = email.message_from_bytes(raw, policy=policy.default)
    matches = []
    for index, part in enumerate(message.walk()):
        if part.is_multipart():
            continue
        payload = part.get_payload(decode=True) or b""
        if not payload.startswith(b"PK"):
            continue
        try:
            book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            header, _, rows = workbook_rows(book, "Operações", PONTE_COLUMNS)
            book.close()
        except Exception:
            continue
        if header is not None:
            matches.append({"name": str(part.get_filename() or f"anexo-{index}.xlsx"),
                            "content": payload, "row_count": len(rows), "header_row": header})
    return message, matches


def _csv_candidates(text):
    candidates = {}
    for delimiter in (";", ",", "\t"):
        try:
            parsed = list(csv_rows(text, delimiter))
        except csv.Error:
            continue
        for role, columns in (("orbe", ORBE_COLUMNS), ("contas", ACCOUNT_COLUMNS), ("aurora", AURORA_COLUMNS)):
            for position, (line, cells) in enumerate(parsed):
                if not matches_header(cells, columns):
                    continue
                rows = [(n, values) for n, values in parsed[position + 1:] if any(v.strip() for v in values)]
                if role == "aurora":
                    label_index = [str(v).strip() for v in cells].index("Lançamento")
                    rows = [(n, values) for n, values in rows
                            if len(values) <= label_index or values[label_index] not in {
                                "SALDO ANTERIOR", "SALDO FINAL", "Os lançamentos estão sujeitos a confirmação."}]
                candidates[role] = {"row_count": len(rows), "header_row": line, "delimiter": delimiter}
                break
    return candidates


def _bank_identity_error(role, text):
    expected = BANK_IDENTITIES.get(role)
    if not expected:
        return None
    if role == "meridiano":
        bank = re.search(r"<BANKID>\s*([^<\r\n]+)", text, re.I)
        if bank and bank.group(1).strip().lstrip("0") != expected["bank_id"]:
            return "O banco informado no OFX difere da identidade Meridiano documentada; confirme a fonte e a conta contábil."
        account = re.search(r"<ACCTID>\s*([^<\r\n]+)", text, re.I)
    else:
        institution = re.search(r"^\s*(Banco[^;\r\n]*)", text, re.M | re.I)
        if institution and not re.match(r"Banco\s+Aurora(?:\s|$)", institution.group(1), re.I):
            return "A instituição informada no CSV difere do Banco Aurora; confirme a fonte e a conta contábil."
        account = re.search(r"\bConta:\s*([^;\r\n]+)", text, re.I)
    if account and re.sub(r"\D", "", account.group(1)) != expected["account"]:
        return "A conta bancária explícita difere da conta documentada para esta fonte; revisão necessária."
    return None


def inspect_bytes(name, content, role=None, file_id=None):
    """Return a JSON-safe manifest entry. Layout errors are data, not exceptions."""
    raw = bytes(content)
    digest = hashlib.sha256(raw).hexdigest()
    original_name = str(name)
    identity = str(file_id) if file_id is not None else "FILE-" + hashlib.sha256(
        (original_name + "\0" + digest).encode("utf-8")).hexdigest()[:24]
    result = {"id": identity, "name": original_name, "sha256": digest, "size": len(raw),
              "role": role if role in ROLES else None, "candidate_roles": [],
              "row_count": 0, "rows": 0, "status": "invalid", "reason": ""}
    if role is not None:
        result["requested_role"] = role

    def finish(status, reason):
        result.update(status=status, reason=reason)
        result["rows"] = result["row_count"]
        result["validation"] = {"status": status, "reason": reason}
        return result

    if role is not None and role not in ROLES:
        return finish("invalid", "Fonte selecionada desconhecida.")
    if not raw or not raw.strip(b"\xef\xbb\xbf \t\r\n"):
        return finish("empty", "Arquivo vazio; nenhum cabeçalho ou registro recebido.")
    candidates = {}
    try:
        if raw.startswith(b"PK"):
            book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            header, _, rows = workbook_rows(book, "Taxas vigentes", RATE_COLUMNS)
            if header is not None:
                candidates["taxas"] = {"row_count": len(rows), "header_row": header, "sheets": list(book.sheetnames)}
            book.close()
        else:
            text, encoding = decode_text(raw)
            result["encoding"] = encoding
            if re.search(r"<OFX\s*>", text, re.I):
                statements = re.findall(r"<STMTTRN>(.*?)</STMTTRN>", text, re.S | re.I)
                # Account section distinguishes a statement from unrelated OFX data;
                # transaction field values are checked per row by the importer.
                if re.search(r"<BANKTRANLIST\s*>", text, re.I) and re.search(r"<BANKACCTFROM\s*>", text, re.I):
                    candidates["meridiano"] = {"row_count": len(statements)}
            elif re.search(r"^(?:MIME-Version|Content-Type):", text, re.M | re.I):
                _, attachments = ponte_attachments(raw)
                if len(attachments) == 1:
                    attachment = attachments[0]
                    candidates["ponte"] = {"row_count": attachment["row_count"],
                                           "header_row": attachment["header_row"],
                                           "attachment_name": attachment["name"]}
                elif len(attachments) > 1:
                    result.update(role="ponte", candidate_roles=["ponte"])
                    return finish("ambiguous", "Mais de um anexo possui a estrutura de operações Ponte; escolha uma fonte sem ambiguidade.")
            else:
                candidates = _csv_candidates(text)
    except (Exception,) as exc:
        # Do not expose parser internals or paths. Invalid bytes remain in manifest.
        return finish("invalid", f"Conteúdo ilegível ou layout incompatível ({type(exc).__name__}).")
    result["candidate_roles"] = sorted(candidates)
    if role is not None:
        if role not in candidates:
            return finish("invalid", "A estrutura do arquivo não corresponde à fonte selecionada.")
        selected = role
    elif len(candidates) > 1:
        return finish("ambiguous", "A estrutura corresponde a mais de uma fonte; selecione a fonte explicitamente.")
    elif not candidates:
        return finish("invalid", "Layout incompatível: cabeçalhos ou estrutura obrigatórios não encontrados.")
    else:
        selected = next(iter(candidates))
    result["role"] = selected
    result.update(candidates[selected])
    identity_error = _bank_identity_error(selected, text) if selected in BANK_IDENTITIES else None
    if identity_error:
        return finish("invalid", identity_error)
    if not result["row_count"]:
        return finish("header_only", "Cabeçalho reconhecido, mas nenhum registro de dados foi recebido.")
    return finish("valid", "Estrutura reconhecida; os valores dos registros serão validados na importação.")


def inspect_file(path, role=None, file_id=None, name=None):
    source = Path(path).expanduser().resolve()
    original_name = str(name) if name is not None else source.name
    try:
        result = inspect_bytes(original_name, source.read_bytes(), role, file_id)
    except OSError:
        result = {"id": str(file_id or "FILE-" + hashlib.sha256(original_name.encode()).hexdigest()[:24]),
                  "name": original_name, "sha256": "", "size": 0, "role": role,
                  "candidate_roles": [], "row_count": 0, "rows": 0,
                  "status": "invalid", "reason": "Não foi possível ler o arquivo apresentado."}
        result["validation"] = {"status": result["status"], "reason": result["reason"]}
    result["path"] = str(source)
    return result


def discover_files(source_dir):
    root = Path(source_dir).expanduser().resolve()
    if not root.is_dir():
        raise ValueError("A pasta de origem não existe.")
    def presented(path):
        name = path.name.casefold()
        suffix = path.suffix.casefold()
        if not path.is_file() or name.startswith((".", "~$")) or name.endswith(("~", ".tmp", ".temp", ".swp")):
            return False
        if suffix in DOCUMENT_EXTENSIONS:
            return False
        if path.stem.casefold() in DOCUMENT_NAMES and suffix in {"", ".txt"}:
            return False
        # Unknown extensions are presented inputs too. Their bytes determine
        # whether they are a valid source or a rejected receipt in the manifest.
        return True
    return [inspect_file(path) for path in sorted(root.iterdir(), key=lambda p: (p.name.casefold(), p.name))
            if presented(path)]
