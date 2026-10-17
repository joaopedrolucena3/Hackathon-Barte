"""Read the supplied exports without rewriting or cleaning their source files."""
from __future__ import annotations

import csv
import email
import hashlib
import io
import re
from datetime import date, datetime
from decimal import Decimal
from email import policy
from pathlib import Path

import openpyxl


SLACK_SOURCE = "Slack #duvidas-hackathon-ai, respostas do financeiro de 16/10/2026"
PRODUCT_COLUMNS = {
    "Débito": "DEBITO", "Crédito à vista": "CREDITO_A_VISTA",
    "Créd. 2-6x": "CREDITO_2_6", "Créd. 7-12x": "CREDITO_7_12",
    "Pix": "PIX", "Antecipação": "ANTECIPACAO", "Spread remessa": "REMESSA",
}
CONFIRMED_EC = {
    "66564512000199": "1043300981", "78576234000179": "1043396662",
    "07752116000126": "1038867654",
}


def decimal_text(value, brazilian=False):
    if value is None or value == "":
        return ""
    original = str(value).strip()
    text = original.replace(".", "").replace(",", ".") if brazilian else original
    try:
        value = Decimal(text)
    except ArithmeticError:
        raise ValueError(f"Número financeiro ilegível: {original!r}.") from None
    if not value.is_finite():
        raise ValueError("Número financeiro não finito.")
    return format(value, "f")


def cnpj_text(value):
    if value is None or value == "":
        return ""
    if isinstance(value, (int, float)):
        value = str(int(value))
    text = re.sub(r"\D", "", str(value))
    return text.zfill(14) if len(text) <= 14 else text


def date_text(value):
    if value is None or value == "":
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).strftime("%Y-%m-%d")
        except ValueError:
            pass
    raise ValueError(f"Data não reconhecida: {text}")


def read_text(path):
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8-sig"), "utf-8-sig"
    except UnicodeDecodeError:
        return raw.decode("cp1252"), "cp1252"


def load_data(source_dir):
    root = Path(source_dir).expanduser().resolve()
    if not root.is_dir():
        raise ValueError("A pasta de origem não existe.")
    data = {k: [] for k in ("merchants", "rates", "accounts", "checkout", "remittances", "bank", "sources", "import_warnings")}
    data["source_dir"] = str(root)
    data["bank_balances"] = {}

    def warning(kind, message, file="", ref="", **extra):
        data["import_warnings"].append({"type": kind, "message": message, "source_file": file, "source_ref": ref, **extra})

    def find(pattern):
        found = sorted(root.glob(pattern))
        if len(found) != 1:
            warning("ARQUIVO_AUSENTE" if not found else "ARQUIVO_AMBIGUO", f"Esperado um arquivo {pattern}; encontrados {len(found)}.", pattern, "arquivo")
            return None
        return found[0]

    def source(path, rows, **extra):
        data["sources"].append({"name": path.name, "path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "rows": rows, **extra})

    path = find("01*comissoes*.csv")
    if path:
        text, encoding = read_text(path)
        reader = csv.DictReader(io.StringIO(text), delimiter=";")
        for row in reader:
            line = reader.line_num
            try:
                data["checkout"].append({
                    "id": row["id_transacao"], "date": date_text(row["data_movimento"]),
                    "ec": row["cod_estabelecimento"], "name": row["nome_estabelecimento"],
                    "type": row["tipo_lancamento"], "modality": row["modalidade"],
                    "installments": int(row["qtd_parcelas"]) if row["qtd_parcelas"] else None,
                    "brand": row["bandeira"], "gross": decimal_text(row["valor_bruto"], True),
                    "applied_rate_pct": decimal_text(row["taxa_aplicada_pct"], True),
                    "reported_fee": decimal_text(row["valor_taxa"], True), "net": decimal_text(row["valor_liquido"], True),
                    "original_id": row["id_referencia"], "original_date": date_text(row["data_venda_original"]),
                    "source_file": path.name, "source_ref": row["id_transacao"], "physical_row": line,
                })
            except (ValueError, KeyError, ArithmeticError) as exc:
                warning("LINHA_INVALIDA", str(exc), path.name, f"L{line}")
        source(path, len(data["checkout"]), encoding=encoding)

    path = find("02*.xlsx")
    if path:
        # In-memory read: read_only workbooks keep the file handle open, which locks it on Windows.
        book = openpyxl.load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
        if "Taxas vigentes" not in book.sheetnames:
            warning("ABA_AUSENTE", "Aba Taxas vigentes não encontrada.", path.name, "workbook")
        else:
            sheet = book["Taxas vigentes"]
            headers = None
            merchants = {}
            for line, cells in enumerate(sheet.values, 1):
                if cells[0] == "Lojista" and "CNPJ" in cells:
                    headers = list(cells)
                    continue
                if not headers or not cells[0] or not cells[1]:
                    continue
                row = dict(zip(headers, cells))
                cnpj = cnpj_text(row.get("CNPJ"))
                ref = f"Taxas vigentes!A{line}:L{line}"
                ec = str(row.get("EC Orbe") or CONFIRMED_EC.get(cnpj, ""))
                active = "INATIVO" not in str(row.get("Obs") or "").upper()
                if cnpj not in merchants:
                    merchants[cnpj] = {"cnpj": cnpj, "name": str(row["Lojista"]), "ec": ec, "active": active,
                                       "candidate": False, "source_file": path.name, "source_ref": ref,
                                       "identity_rule": SLACK_SOURCE if cnpj in CONFIRMED_EC else "CNPJ e EC da tabela vigente"}
                validity = str(row.get("Vigência") or "")
                found = re.search(r"\d{2}/\d{2}/\d{4}", validity)
                start = "" if "até" in validity else date_text(found.group()) if found else ""
                end = date_text(found.group()) if found and "até" in validity else ""
                if not active:
                    continue
                for column, product in PRODUCT_COLUMNS.items():
                    value = row.get(column)
                    default = product == "REMESSA" and (value is None or str(value).strip().lower() == "padrão")
                    if value is None and not default:
                        continue
                    try:
                        pct = Decimal("1.5") if default else (Decimal(str(value).replace("%", "").replace(",", ".")) if isinstance(value, str) else Decimal(str(value)) * 100)
                        if not pct.is_finite():
                            raise ValueError("Taxa não finita.")
                    except (ArithmeticError, ValueError):
                        warning("TAXA_INVALIDA", f"Taxa inválida para {product}.", path.name, ref)
                        continue
                    base = {"id": f"RATE-{cnpj}-{product}-{line}", "cnpj": cnpj, "product": product,
                            "start": start, "end": end, "rate_pct": format(pct, "f"), "source_file": path.name,
                            "source_ref": ref, "rule": SLACK_SOURCE + ": spread padrão 1,5%" if default else "Tabela contratada vigente"}
                    if cnpj == "00520815000100" and product != "REMESSA":
                        base["end"] = "2026-08-31"
                        data["rates"].append(base)
                        data["rates"].append({**base, "id": base["id"] + "-ISENCAO", "start": "2026-09-01", "end": "2026-12-31", "rate_pct": "0", "rule": SLACK_SOURCE + ": Casa isenta MDR e antecipação set–dez; data original da operação"})
                    else:
                        data["rates"].append(base)
            data["merchants"] = list(merchants.values())
            # The old tab is evidence of a candidate identity only. No old rate is imported.
            old_name = "ANTIGO 2025 - nao usar"
            if old_name in book.sheetnames:
                for line, cells in enumerate(book[old_name].values, 1):
                    if cells[0] == "Atlas Idiomas Ltda":
                        data["merchants"].append({"cnpj": cnpj_text(cells[1]), "name": cells[0], "ec": "1076801732",
                            "active": True, "candidate": True, "source_file": path.name, "source_ref": f"{old_name}!A{line}:B{line}",
                            "identity_rule": "Identidade candidata pela aba histórica e nome do relatório Orbe; confirmar cadastro. Taxas antigas não importadas."})
        source(path, sum(1 for r in data["merchants"]), sheets=book.sheetnames)

    path = find("03*.eml")
    if path:
        message = email.message_from_bytes(path.read_bytes(), policy=policy.default)
        data["partner_email"] = {"from": str(message.get("From", "")), "subject": str(message.get("Subject", "")), "date": str(message.get("Date", ""))}
        attachments = [part for part in message.walk() if str(part.get_filename() or "").lower().endswith(".xlsx")]
        if len(attachments) != 1:
            warning("ANEXO_AMBIGUO", f"Esperado um XLSX no EML; encontrados {len(attachments)}.", path.name, "MIME")
        else:
            part = attachments[0]
            payload = part.get_payload(decode=True)
            book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            source_file = f"{path.name}::{part.get_filename()}"
            if "Operações" not in book.sheetnames:
                warning("ABA_AUSENTE", "Aba Operações não encontrada no anexo.", source_file, "workbook")
            else:
                headers = None
                for line, cells in enumerate(book["Operações"].values, 1):
                    if cells[0] == "Nº Operação":
                        headers = list(cells)
                        continue
                    if not headers or not str(cells[0] or "").startswith("PC-"):
                        continue
                    row = dict(zip(headers, cells))
                    try:
                        data["remittances"].append({
                            "id": str(row["Nº Operação"]), "contracted_date": date_text(row["Data Contratação"]),
                            "settlement_date": date_text(row["Data Liquidação"]), "cnpj": cnpj_text(row["CNPJ Cliente"]),
                            "name": row["Cliente"], "currency": row["Moeda"], "foreign_amount": decimal_text(row["Valor ME"]),
                            "client_rate": decimal_text(row["Taxa Cliente"]), "ptax_rate": decimal_text(row["PTAX Ref."]),
                            "gross_brl": decimal_text(row["Valor R$"]), "spread_brl": decimal_text(row["Spread R$"]),
                            "applied_rate_pct": decimal_text(Decimal(str(row["Spread %"])) * 100),
                            "estimated_spread_brl": decimal_text(row["Spread Estimado R$ (D-1)"]), "iof_brl": decimal_text(row["IOF R$"]),
                            "compliance": row["Status Compliance"], "observation": str(row["Observação"] or ""),
                            "source_file": source_file, "source_ref": str(row["Nº Operação"]), "physical_row": line,
                        })
                    except (ValueError, KeyError, ArithmeticError) as exc:
                        warning("LINHA_INVALIDA", str(exc), source_file, f"Operações!L{line}")
            data["sources"].append({"name": source_file, "sha256": hashlib.sha256(payload).hexdigest(), "rows": len(data["remittances"]), "embedded": True})
        source(path, len(data["remittances"]))

    path = find("04*.csv")
    if path:
        text, encoding = read_text(path)
        reader = csv.DictReader(io.StringIO(text), delimiter=";")
        for row in reader:
            data["accounts"].append({"code": row["COD_CONTA"], "name": row["DESCRICAO"], "postable": row["ACEITA_LANCTO"] == "S", "nature": row["NATUREZA"], "source_ref": f"L{reader.line_num}"})
        source(path, len(data["accounts"]), encoding=encoding)

    path = find("05a*.csv")
    if path:
        text, encoding = read_text(path)
        reader = csv.reader(io.StringIO(text), delimiter=";")
        for row in reader:
            line = reader.line_num
            if len(row) >= 5 and row[1] == "SALDO ANTERIOR":
                data["bank_balances"]["AURORA"] = {"opening": decimal_text(row[4], True), "opening_ref": f"L{line}"}
            if len(row) >= 5 and row[1] == "SALDO FINAL":
                data["bank_balances"].setdefault("AURORA", {}).update({"closing": decimal_text(row[4], True), "closing_ref": f"L{line}"})
            if len(row) < 5 or not re.fullmatch(r"\d{2}/\d{2}/\d{4}", row[0]) or not row[3]:
                continue
            try:
                data["bank"].append({"id": f"AURORA-L{line}", "bank": "AURORA", "ref_origem": f"L{line}", "date": date_text(row[0]),
                    "description": row[1], "amount": decimal_text(row[3], True), "balance": decimal_text(row[4], True),
                    "source_file": path.name, "physical_row": line})
            except (ValueError, ArithmeticError) as exc:
                warning("LINHA_INVALIDA", str(exc), path.name, f"L{line}")
        source(path, sum(b["bank"] == "AURORA" for b in data["bank"]), encoding=encoding)

    path = find("05b*.ofx")
    if path:
        text, encoding = read_text(path)
        for match in re.finditer(r"<STMTTRN>(.*?)</STMTTRN>", text, re.S):
            row = dict(re.findall(r"<([^/>]+)>([^<\r\n]*)", match.group(1)))
            line = text[:match.start()].count("\n") + 1
            try:
                data["bank"].append({"id": "MERIDIANO-" + row["FITID"], "bank": "MERIDIANO", "ref_origem": row["FITID"],
                    "date": datetime.strptime(row["DTPOSTED"][:8], "%Y%m%d").date().isoformat(),
                    "description": row["MEMO"], "amount": decimal_text(row["TRNAMT"]), "source_file": path.name, "physical_row": line})
            except (ValueError, KeyError, ArithmeticError) as exc:
                warning("LINHA_INVALIDA", str(exc), path.name, f"L{line}")
        balance = re.search(r"<LEDGERBAL>\s*<BALAMT>([^<\s]+)", text)
        if balance:
            data["bank_balances"]["MERIDIANO"] = {"closing": decimal_text(balance.group(1)), "closing_ref": "LEDGERBAL/BALAMT", "opening": "", "opening_explicit": False}
        source(path, sum(b["bank"] == "MERIDIANO" for b in data["bank"]), encoding=encoding)
    return data
