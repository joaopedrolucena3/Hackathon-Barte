"""Step tables for the data grid: the whole dataset in columnar rows, and exports of the rows the
analyst is looking at. Values leave as the decimal text stored in Parquet; nothing passes through float."""
from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import semantic

FORMATS = {"money": "#,##0.00", "percent": "0.00####", "decimal": "0.######", "number": "#,##0.00##", "integer": "0"}


def table(meta, rows, steps, rows_of):
    dataset = meta["name"]
    columns = semantic.describe(dataset, meta.get("schema", []))
    names = [column["name"] for column in columns]
    return {
        "dataset": dataset, "sha256": meta["sha256"], "row_count": len(rows), "columns": columns,
        "rows": [[row.get(name) for name in names] for row in rows],
        "totals": semantic.totals(dataset, columns, rows),
        "invariants": semantic.invariants(steps, rows_of, only=dataset),
    }


def _text(column, value):
    """CSV cell for Excel pt-BR: decimal comma, dd/mm/aaaa, Sim/Não. Exact string edits only."""
    if value is None or value == "": return ""
    kind = column["kind"]
    if kind in semantic.NUMERIC and not isinstance(value, bool): return str(value).replace(".", ",")
    if kind == "date" and isinstance(value, str) and len(value) == 10 and value[4] == "-": return f"{value[8:10]}/{value[5:7]}/{value[:4]}"
    if isinstance(value, bool): return "Sim" if value else "Não"
    if isinstance(value, (list, dict)): return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _cell(column, value):
    if value is None or value == "": return None
    kind = column["kind"]
    if isinstance(value, bool): return "Sim" if value else "Não"
    if kind in semantic.NUMERIC:
        number = semantic.decimal(value)
        return int(number) if kind == "integer" else number
    if kind == "date" and isinstance(value, str):
        try: return date.fromisoformat(value[:10])
        except ValueError: return value
    if isinstance(value, (list, dict)): return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def export(formato, columns, picked, about):
    """`picked`: [(rid, row)] in the order shown on screen. Returns (filename, media type, bytes)."""
    stem = f"{about['run_id'][:12]}_{about['step_id']}_{about['dataset']}"
    if formato == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream, delimiter=";", lineterminator="\r\n")
        writer.writerow(["Linha no dataset", *[column["label"] for column in columns]])
        for rid, row in picked:
            writer.writerow([rid + 1, *[_text(column, row.get(column["name"])) for column in columns]])
        return stem + ".csv", "text/csv; charset=utf-8", ("﻿" + stream.getvalue()).encode("utf-8")

    book = Workbook()
    sheet = book.active
    sheet.title = about["dataset"][:31]
    sheet.append(["Linha no dataset", *[column["label"] for column in columns]])
    for cell in sheet[1]: cell.font = Font(bold=True)
    for rid, row in picked:
        sheet.append([rid + 1, *[_cell(column, row.get(column["name"])) for column in columns]])
    for index, column in enumerate(columns, start=2):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = max(10, min(42, len(column["label"]) + 4))
        if column["kind"] in FORMATS:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = FORMATS[column["kind"]]
        elif column["kind"] == "date":
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = "DD/MM/YYYY"
    sheet.freeze_panes = "B2"
    info = book.create_sheet("Sobre")
    for key, value in [("Execução", about["run_id"]), ("Passo", f"{about['step_number']}. {about['step_title']} ({about['step_id']})"),
                       ("Tabela", about["dataset"]), ("SHA-256 da tabela", about["sha256"]), ("Linhas exportadas", len(picked)),
                       ("Linhas na tabela", about["row_count"]), ("Filtros e ordenação", about.get("descricao") or "Nenhum"),
                       ("Gerado em (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
                       ("Observação", "Exportado do dataset imutável do passo, na ordem da tela. Valores monetários com duas casas exatas.")]:
        info.append([key, value])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 90
    stream = io.BytesIO()
    book.save(stream)
    return stem + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", stream.getvalue()
